#!/usr/bin/env python3
"""Validate a FaithTrace eval set against docs/dataset/eval_set_schema.md.

Usage:
    python scripts/validate_eval_set.py eval_sets/faithtrace_v1.json
    python scripts/validate_eval_set.py eval_sets/faithtrace_v1.json --check-evidence

Checks: required fields and types, unique ids, enum values, ISO dates
(valid_from <= valid_to), and that every source_doc is listed in
corpus/manifest.json. With --check-evidence, each item's `evidence` quote is
looked up in the source file (needs pdfplumber and openpyxl from
backend/requirements.txt). Exits non-zero on any error.
"""

import argparse
import html
import json
import re
import sys
from datetime import date
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MANIFEST = REPO_ROOT / "corpus" / "manifest.json"

REQUIRED_FIELDS = {
    "id": str,
    "question": str,
    "ground_truth": str,
    "source_docs": list,
    "valid_from": (str, type(None)),
    "valid_to": (str, type(None)),
    "modality": str,
    "difficulty": str,
    "answerable": bool,
    "failure_type": (str, type(None)),
}
MODALITIES = {"text", "table", "chart", "spreadsheet", "mixed"}
DIFFICULTIES = {"easy", "medium", "hard"}
FAILURE_TYPES = {
    "STALE_ANSWER", "WRONG_VERSION", "TABLE_RETRIEVAL_MISS", "CHART_LAYOUT_BLINDNESS",
    "CHUNKING_BOUNDARY_ERROR", "LOW_RECALL_RETRIEVAL", "UNANSWERABLE",
}


def _parse_date(value):
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        return None


def validate_items(items, manifest_files):
    errors = []
    if not isinstance(items, list):
        return ["top-level JSON value must be a list of question objects"]

    seen_ids = set()
    for idx, item in enumerate(items):
        where = f"item[{idx}]"
        if not isinstance(item, dict):
            errors.append(f"{where}: must be an object")
            continue
        where = f"item[{idx}] ({item.get('id', '?')})"

        for field, expected in REQUIRED_FIELDS.items():
            if field not in item:
                errors.append(f"{where}: missing field '{field}'")
            elif not isinstance(item[field], expected):
                errors.append(f"{where}: field '{field}' has wrong type {type(item[field]).__name__}")

        item_id = item.get("id")
        if isinstance(item_id, str):
            if item_id in seen_ids:
                errors.append(f"{where}: duplicate id '{item_id}'")
            seen_ids.add(item_id)

        for field in ("question", "ground_truth"):
            if isinstance(item.get(field), str) and not item[field].strip():
                errors.append(f"{where}: '{field}' is empty")

        if item.get("modality") not in MODALITIES:
            errors.append(f"{where}: invalid modality {item.get('modality')!r}")
        if item.get("difficulty") not in DIFFICULTIES:
            errors.append(f"{where}: invalid difficulty {item.get('difficulty')!r}")
        ft = item.get("failure_type")
        if ft is not None and ft not in FAILURE_TYPES:
            errors.append(f"{where}: invalid failure_type {ft!r}")

        parsed = {}
        for field in ("valid_from", "valid_to"):
            value = item.get(field)
            if value is not None:
                parsed[field] = _parse_date(value)
                if parsed[field] is None:
                    errors.append(f"{where}: '{field}' is not an ISO date (YYYY-MM-DD): {value!r}")
        if parsed.get("valid_from") and parsed.get("valid_to") and parsed["valid_from"] > parsed["valid_to"]:
            errors.append(f"{where}: valid_from is after valid_to")

        docs = item.get("source_docs")
        if isinstance(docs, list):
            for doc in docs:
                if not isinstance(doc, str):
                    errors.append(f"{where}: source_docs entries must be strings")
                elif manifest_files is not None and doc not in manifest_files:
                    errors.append(f"{where}: source_doc '{doc}' not in corpus manifest")
            if item.get("answerable") is True and not docs:
                errors.append(f"{where}: answerable item has no source_docs")
        if item.get("answerable") is False and ft not in (None, "UNANSWERABLE"):
            errors.append(f"{where}: unanswerable item should have failure_type UNANSWERABLE or null")

    return errors


# ---------------------------------------------------------------- evidence checks

def _norm(text):
    text = text.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    text = re.sub(r"\s+", " ", text)
    # Inline HTML tags (e.g. section links) leave a space before punctuation.
    return re.sub(r" ([,.;:])", r"\1", text).strip()


def _html_text(path):
    raw = path.read_text(encoding="utf-8")
    raw = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", raw)
    return _norm(html.unescape(re.sub(r"<[^>]+>", " ", raw)))


def _load_text(path, cache):
    if path in cache:
        return cache[path]
    suffix = path.suffix.lower()
    if suffix in (".html", ".htm"):
        text = _html_text(path)
    elif suffix == ".pdf":
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            text = _norm(" ".join(page.extract_text() or "" for page in pdf.pages))
    elif suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig")
    else:
        text = None
    cache[path] = text
    return text


def _cell_matches(actual, expected):
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return abs(float(actual) - float(expected)) < 1e-6
    return _norm(str(actual)) == _norm(str(expected))


def check_evidence(items, corpus_dir):
    """Return (checked, errors, visual_only) for answerable items with evidence."""
    errors, visual_only, cache, checked = [], [], {}, 0
    workbooks = {}
    for item in items:
        if not item.get("answerable"):
            continue
        ev = item.get("evidence")
        where = item.get("id", "?")
        if not isinstance(ev, dict) or not (ev.get("quote") or ev.get("cells")):
            errors.append(f"{where}: answerable item has no evidence quote")
            continue
        quotes = ev.get("quote")
        quotes = quotes if isinstance(quotes, list) else ([quotes] if quotes else [])
        for doc in item.get("source_docs", []):
            path = corpus_dir / doc
            if not path.exists():
                errors.append(f"{where}: source file missing: {path}")
                continue
            if path.suffix.lower() == ".xlsx":
                import openpyxl
                if path not in workbooks:
                    workbooks[path] = openpyxl.load_workbook(str(path), read_only=False)
                sheet = re.search(r"sheet '([^']+)'", ev.get("location") or "")
                if not sheet or sheet.group(1) not in workbooks[path].sheetnames:
                    errors.append(f"{where}: evidence location must name an existing sheet")
                    continue
                ws = workbooks[path][sheet.group(1)]
                for ref, expected in (ev.get("cells") or {}).items():
                    if not _cell_matches(ws[ref].value, expected):
                        errors.append(f"{where}: {sheet.group(1)}!{ref} is {ws[ref].value!r}, expected {expected!r}")
                checked += 1
                continue
            text = _load_text(path, cache)
            if text is None:
                errors.append(f"{where}: cannot read text from {doc}")
                continue
            haystack = _norm(text)
            missing = [q for q in quotes if _norm(q) not in haystack]
            if missing and "chart image" in (ev.get("location") or ""):
                visual_only.append(f"{where}: {missing} (in chart image only; verify visually)")
            elif missing:
                errors.append(f"{where}: quote not found in {doc}: {missing}")
            checked += 1
    return checked, errors, visual_only


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("eval_set", type=Path)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--check-evidence", action="store_true",
                        help="re-open each source file and confirm the evidence quote/cells exist")
    args = parser.parse_args()

    try:
        items = json.loads(args.eval_set.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        print(f"ERROR: cannot load {args.eval_set}: {exc}")
        return 1

    manifest_files = None
    if args.manifest.exists():
        manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
        manifest_files = {entry["filename"] for entry in manifest.get("documents", [])}
    else:
        print(f"ERROR: manifest not found at {args.manifest}")
        return 1

    errors = validate_items(items, manifest_files)
    visual_only = []
    if args.check_evidence and not errors:
        checked, ev_errors, visual_only = check_evidence(items, args.manifest.parent)
        errors.extend(ev_errors)
        print(f"Evidence checked for {checked} answerable items.")
        for note in visual_only:
            print(f"NOTE: {note}")

    if errors:
        for err in errors:
            print(f"ERROR: {err}")
        print(f"FAILED: {len(errors)} error(s) in {args.eval_set}")
        return 1
    print(f"OK: {args.eval_set} ({len(items)} items)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
